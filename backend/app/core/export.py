"""Export: trades/executions CSV, a full Excel workbook, and a backup zip."""

from __future__ import annotations

import csv
import io
import sqlite3
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import settings
from app.models import (
    Account,
    DailySession,
    Execution,
    Instrument,
    MistakeFlag,
    Tag,
    Trade,
)

TRADE_COLUMNS = [
    "id", "account", "symbol", "direction", "trading_date", "open_time", "close_time",
    "status", "max_position_size", "avg_entry", "avg_exit", "points", "gross_pnl",
    "fees", "net_pnl", "currency", "initial_stop", "planned_target", "planned_risk_amount",
    "r_multiple", "holding_time_seconds", "scenario", "self_rating", "pnl_only", "source",
    "style",
]

EXEC_COLUMNS = [
    "id", "account_id", "instrument_id", "side", "quantity", "price", "timestamp",
    "fees", "source", "external_id", "time_known", "trade_id",
]


def _trade_row(db: Session, t: Trade) -> list:
    return [
        t.id, t.account.name if t.account else "", t.instrument.symbol_root if t.instrument else "",
        t.direction, t.trading_date, t.open_time, t.close_time, t.status, t.max_position_size,
        t.avg_entry, t.avg_exit, t.points, t.gross_pnl, t.fees, t.net_pnl, t.currency,
        t.initial_stop, t.planned_target, t.planned_risk_amount, t.r_multiple,
        t.holding_time_seconds, t.scenario, t.self_rating, t.pnl_only, t.source,
        t.style,
    ]


def trades_csv(db: Session) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(TRADE_COLUMNS)
    for t in db.scalars(select(Trade).order_by(Trade.open_time)):
        w.writerow([_fmt(v) for v in _trade_row(db, t)])
    return buf.getvalue()


def executions_csv(db: Session) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(EXEC_COLUMNS)
    for e in db.scalars(select(Execution).order_by(Execution.timestamp)):
        w.writerow([_fmt(getattr(e, c)) for c in EXEC_COLUMNS])
    return buf.getvalue()


def _fmt(v):
    if isinstance(v, datetime):
        return v.isoformat()
    return "" if v is None else str(v)


def workbook_xlsx(db: Session) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "trades"
    ws.append(TRADE_COLUMNS)
    for t in db.scalars(select(Trade).order_by(Trade.open_time)):
        ws.append([_fmt(v) for v in _trade_row(db, t)])

    we = wb.create_sheet("executions")
    we.append(EXEC_COLUMNS)
    for e in db.scalars(select(Execution).order_by(Execution.timestamp)):
        we.append([_fmt(getattr(e, c)) for c in EXEC_COLUMNS])

    wt = wb.create_sheet("tags")
    wt.append(["id", "name", "category", "archived"])
    for tag in db.scalars(select(Tag)):
        wt.append([tag.id, tag.name, tag.category, tag.archived])

    wd = wb.create_sheet("daily_sessions")
    wd.append(["id", "account_id", "instrument_id", "trading_date", "actual_day_type",
               "actual_opening_type", "followed_plan", "mood", "sleep"])
    for d in db.scalars(select(DailySession)):
        wd.append([d.id, d.account_id, d.instrument_id, _fmt(d.trading_date), d.actual_day_type,
                   d.actual_opening_type, d.followed_plan, d.mood, d.sleep])

    wm = wb.create_sheet("mistake_flags")
    wm.append(["id", "trade_id", "account_id", "session_date", "rule_code", "source",
               "dismissed", "details"])
    for m in db.scalars(select(MistakeFlag)):
        wm.append([m.id, m.trade_id, m.account_id, _fmt(m.session_date), m.rule_code,
                   m.source, m.dismissed, m.details])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def backup_zip(db: Session) -> bytes:
    """A consistent SQLite snapshot (VACUUM INTO) plus the attachments folder."""
    out = io.BytesIO()
    with tempfile.TemporaryDirectory() as tmp:
        snapshot = Path(tmp) / "journal.db"
        con = sqlite3.connect(str(settings.db_path))
        try:
            con.execute("VACUUM INTO ?", (str(snapshot),))
        finally:
            con.close()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.write(snapshot, "journal.db")
            adir = settings.attachments_dir
            if adir.exists():
                for f in adir.rglob("*"):
                    if f.is_file():
                        zf.write(f, f"attachments/{f.relative_to(adir).as_posix()}")
    return out.getvalue()


def backup_filename() -> str:
    return f"trade_journal_backup_{datetime.now(timezone.utc):%Y%m%d_%H%M%S}.zip"


def restore_zip(content: bytes) -> dict:
    """Replace the database and attachments from a backup zip."""
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        names = zf.namelist()
        if "journal.db" not in names:
            return {"status": "invalid", "message": "journal.db missing from backup"}
        settings.ensure_dirs()
        # Replace the database file.
        data = zf.read("journal.db")
        settings.db_path.write_bytes(data)
        # Clear WAL/SHM side files so the restored DB is authoritative.
        for side in ("-wal", "-shm"):
            p = Path(str(settings.db_path) + side)
            if p.exists():
                p.unlink()
        # Restore attachments.
        for name in names:
            if name.startswith("attachments/") and not name.endswith("/"):
                target = settings.attachments_dir / name[len("attachments/"):]
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(zf.read(name))
    return {"status": "restored"}
